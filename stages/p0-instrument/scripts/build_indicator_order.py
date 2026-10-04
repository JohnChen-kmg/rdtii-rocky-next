"""Step 0 — emit output/indicator_order.yaml, the one ordered list of indicator IDs.

Source: the host's finale template, sheet 'Indicator Reference' (62 IDs in the host's own
order, with names, pillar weights and exception notes, plus the host's "five mapping traps"
block under the list). Every other artefact keys off this file: the Tier C generator decides
which IDs to emit from it, the mapper groups by pillar with it, and the validator proves every
in-scope ID is defined exactly once. Order is never derived from the numeric value of an ID.

It also writes the coverage marks of decisions D13 and D14 onto every entry, from
scripts/data/coverage.yaml: whether the tool automates the indicator, the reason a manual one is
manual, and the nature of its answer. The mapping stage's automated set is the entries marked
`coverage: automated`.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import yaml

from indicator_ids import normalize, pillar_of
from rdtii_examples import DATA, REPO, TEMPLATE_FINAL, load_methodology

OUT = REPO / "output" / "indicator_order.yaml"
COVERAGE = DATA / "coverage.yaml"
SHEET = "Indicator Reference"

# The host's non-regulatory indicators note (ESCAP-RDTII-2.1 Non-regulatory indicators.pdf):
# derived from external databases and internet research, not extracted by the tool.
NON_REGULATORY = ["1.1", "1.2", "1.3", "2.4", "4.4", "4.7", "4.8", "5.6", "6.5", "9.2",
                  "12.10", "12.11", "12.12", "12.13"]

# Internal Guide p.8: sub-pillars that focus on enforcement and practice, where secondary
# sources may be used as evidence alongside (or instead of) legal text.
PRACTICE_BASED = {
    "3.4": "Screening of investment: the top score needs a case where screening blocked an investment (Internal Guide p.8; Methodology sheet 3.4)",
    "5.3": "Government shares in telecom companies: ownership facts from SOE lists, market reports and company records (Internal Guide pp.8, 12)",
    "9.1": "Blocking or filtering commercial web content: blocking practice can be evidenced by official announcements and secondary reports (Internal Guide pp.8, 14)",
}

OUT_OF_SCOPE_REASON = {
    "6.5": ("Non-regulatory: participation in agreements with binding data-transfer commitments is checked "
            "from treaty texts, not national legislation. The methodology sheet has no criteria row for it, "
            "its header says gaps in the ID sequence are intentional and the tool is not required to extract "
            "non-regulatory indicators, and the template note says it is not present in the database extract."),
}


def coverage_marks(iid: str, status: str, cov: dict) -> dict:
    """The coverage class, the manual-check reason and the nature of the answer (decisions D13, D14).

    coverage         automated | manual | excluded
    coverage_reason  a key of the mapping stage's reason categories, for a manual indicator:
                     "practice" when no legal instrument states the answer in any economy,
                     otherwise "scope"
    answer_nature    the coverage register's nature codes, operative code first
    """
    if status != "in_scope":
        return {"coverage": "excluded", "coverage_reason": None, "answer_nature": None}
    nature = {str(k): v for k, v in cov["answer_nature"].items()}.get(iid)
    if not nature:
        raise SystemExit(f"coverage.yaml: no answer_nature for {iid}")
    unknown = [c for c in nature if c not in cov["nature_codes"]]
    if unknown:
        raise SystemExit(f"coverage.yaml: {iid} uses unknown nature codes {unknown}")
    automated = pillar_of(iid) in cov["automated_pillars"] or iid in [str(x) for x in cov["automated_indicators"]]
    if automated:
        return {"coverage": "automated", "coverage_reason": None, "answer_nature": list(nature)}
    reason = "practice" if nature[0] in cov["outside_any_law_database"] else "scope"
    return {"coverage": "manual", "coverage_reason": reason, "answer_nature": list(nature)}


def main() -> None:
    wb = openpyxl.load_workbook(TEMPLATE_FINAL, data_only=True)
    ws = wb[SHEET]
    meth = load_methodology()
    cov = yaml.safe_load(COVERAGE.read_text(encoding="utf-8"))

    host_note = str(ws.cell(2, 1).value or "").strip()
    entries, pillar_label = [], None
    trap_rows: list[dict] = []
    in_trap_block = False
    for r in range(4, ws.max_row + 1):
        a, b, c, d, e = (ws.cell(r, k).value for k in range(1, 6))
        if in_trap_block:
            if a:
                trap_rows.append({"template_row": r, "text": str(a).strip()})
            continue
        if a and str(a).strip().lower().startswith("the five mapping traps"):
            in_trap_block = True
            trap_rows.append({"template_row": r, "text": str(a).strip()})
            continue
        if a and b is None:
            pillar_label = str(a).strip()
            continue
        if b is None:
            continue
        iid = normalize(str(b))
        if iid != str(b).strip():
            raise SystemExit(f"row {r}: template ID {b!r} is not canonical text")
        status = "out_of_scope" if iid in OUT_OF_SCOPE_REASON else "in_scope"
        if status == "in_scope" and iid not in meth:
            raise SystemExit(f"row {r}: {iid} has no methodology row but is not declared out of scope")
        entries.append({
            "id": iid,
            "pillar": pillar_of(iid),
            "pillar_label": pillar_label,
            "name": str(c or "").strip(),
            "weight_in_pillar": d,
            "host_note": (str(e).strip() if e else None),
            "template_row": r,
            "methodology_row": meth[iid]["row"] if iid in meth else None,
            "status": status,
            "evidence": "practice" if iid in PRACTICE_BASED else "legal",
            **coverage_marks(iid, status, cov),
        })

    ids = [x["id"] for x in entries]
    if len(ids) != len(set(ids)):
        raise SystemExit("duplicate IDs in the Indicator Reference sheet")
    for iid in PRACTICE_BASED:      # the host's practice-based list must come out with the practice reason
        e = next(x for x in entries if x["id"] == iid)
        if e["coverage"] == "manual" and e["coverage_reason"] != "practice":
            raise SystemExit(f"coverage.yaml: {iid} is practice-based (Internal Guide p.8) but its nature code gives '{e['coverage_reason']}'")
    doc = {
        "source": {
            "workbook": "reference/OUTPUT_TEMPLATE_FINAL_ROUND.xlsx",
            "sheet": SHEET,
            "host_note": host_note,
        },
        "listed": len(entries),
        "in_scope": sum(1 for x in entries if x["status"] == "in_scope"),
        "out_of_scope": {k: v for k, v in OUT_OF_SCOPE_REASON.items() if k in ids},
        "non_regulatory_indicators": NON_REGULATORY,
        "practice_based": PRACTICE_BASED,
        "host_mapping_traps": trap_rows,
        # Decisions D13 and D14: what the tool automates, why a manual indicator is manual, and what
        # kind of thing each answer is. Restated from scripts/data/coverage.yaml.
        "coverage_marks": {
            "register": cov["register"],
            "classes": cov["classes"],
            "reasons": cov["reasons"],
            "nature_codes": cov["nature_codes"],
            "counts": {c: sum(1 for x in entries if x["coverage"] == c) for c in ("automated", "manual", "excluded")},
            "economy_overrides": cov["economy_overrides"],
        },
        "indicators": entries,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="\n") as f:
        f.write("# Generated by scripts/build_indicator_order.py from the host template's\n"
                "# 'Indicator Reference' sheet. The one ordered list of indicator IDs; do not hand-edit.\n")
        yaml.safe_dump(doc, f, allow_unicode=True, sort_keys=False, width=110)
    print(f"wrote {OUT.relative_to(REPO)}: {doc['listed']} listed, {doc['in_scope']} in scope, "
          f"{len(trap_rows)} host trap rows; coverage {doc['coverage_marks']['counts']}")


if __name__ == "__main__":
    main()
