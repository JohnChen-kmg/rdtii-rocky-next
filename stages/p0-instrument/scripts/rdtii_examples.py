"""Shared parser for the RDTII host workbooks (all pillars, decimal indicator IDs).

Encodes the verified parsing hazards:
  - country sheets ONLY (the Round 1 'Consolidated' sheet is lossy — never read it)
  - Indicator_ID cells may be merged across rows (MY B49:B51, SG B38:B39):
    forward-fill the last seen indicator ID; any other text in column B (a pillar or
    section header) RESETS the fill so header blocks are never mis-attributed
  - IDs arrive as floats for two-level IDs (7.3 -> 7.3000000001 style) and as text for
    three-level IDs (12.4.1). Both go through indicator_ids.normalize(); IDs are never
    compared as numbers (4.01 and 4.1 are different indicators)
  - strikethrough formatting = deleted content — surfaced as a flag, callers drop

Finale change (W3, 2026-09-13): Round 1 kept pillars 6-7 only and rewrote IDs as P6-I1.
This version keeps every pillar and emits decimal text IDs.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from indicator_ids import BadIndicatorId, normalize

SCRIPTS = Path(__file__).resolve().parent
DATA = SCRIPTS / "data"


def _stage_folder() -> Path:
    """The instrument stage folder (the one holding output/, examples/ and reference/).

    Two layouts are supported:
      finale repo           stages/p0-instrument/scripts/  -> the stage is the parent folder
      instrument workspace  <workspace>/code/scripts/      -> the stage is <workspace>/instrument/
    """
    for stage in (SCRIPTS.parent, SCRIPTS.parent.parent / "instrument"):
        if (stage / "output").is_dir():
            return stage
    raise SystemExit(f"cannot find the instrument stage folder (an output/ folder) near {SCRIPTS}")


REPO = _stage_folder()
ROUND1 = REPO / "examples" / "Round1_Baseline_Database.xlsx"
ROUND2 = REPO / "examples" / "Round2_Methodology_and_Examples.xlsx"
TEMPLATE_FINAL = REPO / "reference" / "OUTPUT_TEMPLATE_FINAL_ROUND.xlsx"
METHODOLOGY_SHEET = "RDTII 2.1 Methodology"

ROUND1_SHEETS = ["Australia", "Malaysia", "Singapore"]
ROUND2_SHEETS = ["China", "India", "Indonesia", "Lao PDR", "Mongolia",
                 "Russian Federation", "Thailand"]

ECONOMY_CODE = {
    "Australia": "AU", "Malaysia": "MY", "Singapore": "SG",
    "China": "CN", "India": "IN", "Indonesia": "ID", "Lao PDR": "LA",
    "Mongolia": "MN", "Russian Federation": "RU", "Thailand": "TH",
}

_ART_RE = re.compile(
    r"(?:Section|Sections|Article|Articles|Schedule|Part|Rule|Regulation|"
    r"s\.|ss\.|Art\.|§)\s*\d+[A-Za-z]{0,2}(?:\s*\(\d+\))?(?:\s*\([a-z]\))?"
)


def _clean(v):
    if v is None:
        return None
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def _urls(cells):
    out = []
    for v in cells:
        if not v:
            continue
        # allow parentheses inside URLs (e.g. ...(PDPA)%20Act%202010.pdf);
        # strip only trailing sentence punctuation
        for u in re.findall(r"https?://[^\s;,\]]+", str(v)):
            u = u.rstrip(".,;)")
            if u not in out:
                out.append(u)
    return out


def _indicator_cell(v) -> str | None:
    """Return the decimal ID if the column-B cell holds one, else None."""
    if v is None:
        return None
    try:
        return normalize(v)
    except BadIndicatorId:
        return None


def gold_id(source: str, sheet: str, row: int) -> str:
    """Stable row ID: r1-my-053 (Round 1 workbook) · r2-in-095 (Round 2 workbook)."""
    return f"{'r1' if source == 'round1' else 'r2'}-{ECONOMY_CODE[sheet].lower()}-{row:03d}"


def articles_mentioned(text: str | None) -> list[str]:
    if not text:
        return []
    seen, out = set(), []
    for m in _ART_RE.findall(text):
        key = re.sub(r"\s+", " ", m).strip()
        if key.lower() not in seen:
            seen.add(key.lower())
            out.append(key)
    return out


def _extra_columns(ws) -> dict:
    """Header-driven lookup of the columns after References (L onwards), which differ per sheet.

    Verified 2026-09-13: 'Note' sits in L or M; Indonesia adds 'Identify types of update';
    Lao PDR and Russian Federation add a government-verification question; Thailand adds
    'Data verification feedback' (Correct / Not correct), its full text, and 'Actions'.
    """
    hdr = {c: re.sub(r"\s+", " ", str(ws.cell(1, c).value or "")).strip().lower()
           for c in range(12, ws.max_column + 1)}
    pick = lambda pred: [c for c, h in hdr.items() if h and pred(h)]
    return {
        "note": pick(lambda h: h.startswith("note")),
        "update_type": pick(lambda h: h.startswith("identify types of update")),
        "question": pick(lambda h: "specific question" in h),
        "feedback": pick(lambda h: h == "data verification feedback"),
        "feedback_full": pick(lambda h: h.startswith("data verification feedback (full)")),
        "action": pick(lambda h: h == "actions"),
    }


def _first(ws, r: int, cols: list[int]):
    for c in cols:
        v = _clean(ws.cell(r, c).value)
        if v:
            return v
    return None


def parse_workbook(path: Path, sheets: list[str], source: str,
                   pillars: set[int] | None = None) -> list[dict]:
    """Every coded row in the country sheets. `pillars` optionally filters (None = all)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for sheet in sheets:
        ws = wb[sheet]
        extra = _extra_columns(ws)
        cur = None
        for r in range(2, ws.max_row + 1):
            raw = ws.cell(r, 2).value
            iid = _indicator_cell(raw)
            if iid:
                cur = iid
            elif raw is not None and str(raw).strip():  # header text resets the forward-fill
                cur = None
            if not cur:
                continue
            pillar = int(cur.split(".")[0])
            if pillars is not None and pillar not in pillars:
                continue
            act = _clean(ws.cell(r, 4).value)
            imp = _clean(ws.cell(r, 6).value)
            if act is None and imp is None:
                continue
            struck = any(
                ws.cell(r, c).value is not None
                and ws.cell(r, c).font and ws.cell(r, c).font.strike
                for c in range(2, 14)
            )
            rows.append({
                "source": source,
                "sheet": sheet,
                "economy": ECONOMY_CODE[sheet],
                "row": r,
                "gold_id": gold_id(source, sheet, r),
                "indicator": cur,
                "pillar": pillar,
                "raw_score": ws.cell(r, 3).value,
                "law": act,
                "coverage": _clean(ws.cell(r, 5).value),
                "impact": imp,
                "timeframe": _clean(ws.cell(r, 7).value),
                "urls": _urls([ws.cell(r, c).value for c in (8, 9, 10, 11)]),
                "note": _first(ws, r, extra["note"]),
                "update_type": _first(ws, r, extra["update_type"]),
                "host_verification": {k: v for k, v in {
                    "question": _first(ws, r, extra["question"]),
                    "feedback": _first(ws, r, extra["feedback"]),
                    "feedback_full": _first(ws, r, extra["feedback_full"]),
                    "action": _first(ws, r, extra["action"]),
                }.items() if v} or None,
                "strikethrough": struck,
            })
    return rows


def load_all(pillars: set[int] | None = None) -> list[dict]:
    return (parse_workbook(ROUND1, ROUND1_SHEETS, "round1", pillars)
            + parse_workbook(ROUND2, ROUND2_SHEETS, "round2", pillars))


def load_methodology(path: Path = ROUND1) -> dict[str, dict]:
    """The host 'RDTII 2.1 Methodology' sheet keyed by decimal ID.

    The finale template's Indicator Reference names this sheet in the Round 1 database as
    its source; the Round 2 database carries an identical copy (checked 2026-09-13).
    """
    ws = openpyxl.load_workbook(path, data_only=True)[METHODOLOGY_SHEET]
    out: dict[str, dict] = {}
    for r in range(2, ws.max_row + 1):
        iid = _indicator_cell(ws.cell(r, 2).value)
        if not iid:
            continue
        out[iid] = {
            "row": r,
            "category": str(ws.cell(r, 3).value or ""),
            "criteria": str(ws.cell(r, 4).value or ""),
            "possible_scores": str(ws.cell(r, 5).value or ""),
        }
    return out


def score_values(possible_scores: str) -> list[float]:
    """'1\\n0.50\\n0' -> [1.0, 0.5, 0.0]."""
    return [float(x) for x in re.findall(r"\d+(?:\.\d+)?", possible_scores or "")]


def category_first_line(category: str) -> str:
    return re.sub(r"\s+", " ", str(category).split("\n")[0]).strip()


def trim_sentence(text: str | None, limit: int = 650) -> str | None:
    """Trim to <= limit chars, cutting at the last sentence boundary."""
    if text is None:
        return None
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= limit:
        return text
    cut = text[:limit]
    m = max(cut.rfind(". "), cut.rfind("; "))
    return (cut[: m + 1] if m > limit // 2 else cut.rstrip() + "…").strip()
