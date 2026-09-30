"""The submission folder holds ONE file per economy, whatever number of passes produced it.

The host reads one records file, one workbook and one audit page per economy. This run does not
naturally produce that shape: Timor-Leste is mapped in two arms, so it had two of each. These tests
pin the merge and, more importantly, the guard that catches a half-merged country -- a workbook
copied from one arm looks completely normal and is simply missing 52 indicators.
"""
from __future__ import annotations

import csv
import json

import pytest

from src.p3map.output.package import CSV_COLUMNS, build, merge_csv, merge_json


def _row(econ="Timor-Leste", ind="6.1", law="An Act 2020", sec="Art. 1", snip="text"):
    return {c: "" for c in CSV_COLUMNS} | {"Economy": econ, "Indicator ID": ind, "Law Name": law,
                                           "Article / Section": sec, "Verbatim Snippet": snip,
                                           "Language of Source": "Portuguese"}


def _write_csv(p, rows):
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CSV_COLUMNS)
        w.writeheader()
        w.writerows(rows)
    return p


def _arm(root, econ, rows, indicators_fired, with_workbook=True):
    """A run directory with a records file, a per-law JSON, verdicts, and optionally a workbook."""
    d = root
    _write_csv(d / "submission" / f"records_{econ}.csv", rows)
    (d / "submission" / f"records_{econ}.json").write_text(json.dumps({
        "contract_version": "0.3.0", "economy": econ, "economy_name": "Timor-Leste",
        "laws": [{"law_name": "An Act 2020", "law_number": "1/2020",
                  "source_url": "https://x.gov/a.pdf",
                  "provisions": [{"Indicator ID": r["Indicator ID"]} for r in rows]}]}),
        encoding="utf-8")
    (d / "map").mkdir(parents=True, exist_ok=True)
    with (d / "map" / f"verdicts_{econ}.jsonl").open("w", encoding="utf-8") as f:
        for ind in indicators_fired:
            f.write(json.dumps({"provision_id": f"p-{ind}", "trap_checks": {},
                                "verdicts": [{"indicator": ind, "applies": True}]}) + "\n")
    if with_workbook:
        _fake_workbook(d / "results" / f"RDTII_P3_results_{econ}.xlsx", indicators_fired)
    return d


def _fake_workbook(path, indicators):
    from openpyxl import Workbook
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "All fires"
    ws.cell(row=1, column=1, value="• banner")
    ws.cell(row=2, column=1, value="Indicator")
    for i, ind in enumerate(indicators, start=3):
        ws.cell(row=i, column=1, value=ind)
    wb.save(path)


# --------------------------------------------------------------------------- merging


def test_the_arms_records_are_concatenated():
    """The arms cover disjoint indicator sets — 9 and 52 on this run, zero overlap — so this is a
    concatenation, not a union with dedupe. A dedupe would silently drop a real row the day they do
    overlap."""
    import tempfile
    from pathlib import Path
    with tempfile.TemporaryDirectory() as t:
        a = _write_csv(Path(t) / "a.csv", [_row(ind="6.1"), _row(ind="7.3")])
        b = _write_csv(Path(t) / "b.csv", [_row(ind="3.5")])
        rows = merge_csv([a, b])
    assert [r["Indicator ID"] for r in rows] == ["6.1", "7.3", "3.5"]


def test_a_records_file_missing_a_host_column_is_refused(tmp_path):
    p = tmp_path / "bad.csv"
    with p.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Economy", "Indicator ID"])
        w.writeheader()
        w.writerow({"Economy": "Timor-Leste", "Indicator ID": "6.1"})
    with pytest.raises(ValueError, match="missing host columns"):
        merge_csv([p])


def test_a_law_cited_in_both_arms_becomes_one_entry(tmp_path):
    """The host's shape is one entry per law, not per pass."""
    def js(p, inds):
        p.write_text(json.dumps({"economy": "TL", "laws": [
            {"law_name": "An Act 2020", "law_number": "1/2020", "source_url": "u",
             "provisions": [{"Indicator ID": i} for i in inds]}]}), encoding="utf-8")
        return p
    a = js(tmp_path / "a.json", ["6.1"])
    b = js(tmp_path / "b.json", ["3.5"])
    merged = merge_json([a, b])
    assert len(merged["laws"]) == 1
    assert [p["Indicator ID"] for p in merged["laws"][0]["provisions"]] == ["6.1", "3.5"]
    assert merged["economy"] == "TL", "the top-level fields survive"


def test_two_different_laws_stay_separate(tmp_path):
    a = tmp_path / "a.json"
    a.write_text(json.dumps({"economy": "TL", "laws": [
        {"law_name": "Act A", "law_number": "1", "source_url": "u1", "provisions": []}]}),
        encoding="utf-8")
    b = tmp_path / "b.json"
    b.write_text(json.dumps({"economy": "TL", "laws": [
        {"law_name": "Act B", "law_number": "2", "source_url": "u2", "provisions": []}]}),
        encoding="utf-8")
    assert len(merge_json([a, b])["laws"]) == 2


# --------------------------------------------------------------------------- the folder


def test_one_records_file_and_one_workbook_per_economy(tmp_path):
    p1 = _arm(tmp_path / "out", "TL", [_row(ind="6.1")], ["6.1"])
    p2 = _arm(tmp_path / "out_tl52", "TL", [_row(ind="3.5")], ["3.5"], with_workbook=False)
    out = tmp_path / "submission"
    rep = build({"TL": [p1, p2]}, out)
    assert sorted(p.name for p in out.glob("records_TL.*")) == ["records_TL.csv",
                                                               "records_TL.json"]
    assert (out / "RDTII_P3_results_TL.xlsx").exists()
    assert rep["economies"]["TL"]["rows"] == 2
    assert rep["economies"]["TL"]["indicators"] == 2


def test_a_half_merged_workbook_is_reported_not_shipped_silently(tmp_path):
    """The failure this guard exists for: a workbook copied from one arm looks entirely normal and
    is simply missing the other arm's indicators. Nothing else in the folder would show it."""
    p1 = _arm(tmp_path / "out", "TL", [_row(ind="6.1")], ["6.1"])          # workbook: 6.1 only
    p2 = _arm(tmp_path / "out_tl52", "TL", [_row(ind="3.5")], ["3.5"], with_workbook=False)
    rep = build({"TL": [p1, p2]}, tmp_path / "submission")
    assert rep["warnings"], "a missing arm must be reported"
    assert "3.5" in rep["warnings"][0]
    assert "--extra-dirs" in rep["warnings"][0], "and the message must say how to fix it"


def test_a_properly_merged_workbook_raises_no_warning(tmp_path):
    p1 = _arm(tmp_path / "out", "TL", [_row(ind="6.1")], ["6.1", "3.5"])   # merged workbook
    p2 = _arm(tmp_path / "out_tl52", "TL", [_row(ind="3.5")], ["3.5"], with_workbook=False)
    rep = build({"TL": [p1, p2]}, tmp_path / "submission")
    assert rep["warnings"] == []


def test_a_single_arm_economy_needs_no_special_handling(tmp_path):
    p = _arm(tmp_path / "out", "SG", [_row(econ="Singapore", ind="6.1")], ["6.1"])
    rep = build({"SG": [p]}, tmp_path / "submission")
    assert rep["economies"]["SG"]["rows"] == 1
    assert rep["warnings"] == []


def test_the_report_counts_scored_rows_separately(tmp_path):
    rows = [_row(ind="6.1"), _row(ind="6.2", snip="No provision found")]
    p = _arm(tmp_path / "out", "TL", rows, ["6.1"])
    rep = build({"TL": [p]}, tmp_path / "submission")
    assert rep["economies"]["TL"]["rows"] == 2
    assert rep["economies"]["TL"]["scored"] == 1


def test_the_package_report_is_written_to_the_folder(tmp_path):
    p = _arm(tmp_path / "out", "TL", [_row()], ["6.1"])
    out = tmp_path / "submission"
    build({"TL": [p]}, out)
    saved = json.loads((out / "package_report.json").read_text(encoding="utf-8"))
    assert saved["total_rows"] == 1
