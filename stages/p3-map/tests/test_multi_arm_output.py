"""One workbook and one audit page per country, even when the country was mapped in two arms.

Timor-Leste is mapped twice: pillars 6 and 7 in `out/` and the other 52 indicators in `out_tl52/`.
That gave it two workbooks, two audit pages and two records files while every other economy had one,
and left a reviewer to reconcile them. The submission needs one file per economy regardless.

The trap these tests exist for: a provision can be judged in BOTH arms, because the arms cover
different indicators but the same article can be a candidate under one of pillar 6's and one of
pillar 3's. Fires are per (provision x indicator) and all of them belong in the workbook; the
provision-LEVEL surfaces -- the mapped count, QA NotInForce, QA Errors -- must not double-count.
"""
from __future__ import annotations

import dataclasses
import json

import pytest

from config.settings import SETTINGS
from src.p3map.output import excel_export as ex
from src.p3map.verify import audit_view as av


def _verdict(pid, indicator, applies=True, in_force=True, quote="shall be stored domestically"):
    return {"provision_id": pid, "economy": "TL", "doc_id": "d1", "law_name": "An Act 2020",
            "article_section": "Art. 1", "trap_checks": {"provision_in_force": in_force},
            "verdicts": [{"indicator": indicator, "applies": applies, "score_hint": "1",
                          "coverage": "Horizontal", "confidence": 0.9, "quote_grounded_ws": True,
                          "verbatim_quote": quote, "rationale": "branch"}]}


@pytest.fixture
def arms(tmp_path, monkeypatch):
    """Two arm directories sharing one provision, plus a minimal handoff2."""
    primary, second = tmp_path / "out", tmp_path / "out_tl52"
    for d in (primary, second):
        for sub in ("map", "verify", "audit", "results"):
            (d / sub).mkdir(parents=True)

    # p-both is judged in both arms; p-a and p-b only in one each
    (primary / "map" / "verdicts_TL.jsonl").write_text(
        json.dumps(_verdict("p-both", "6.1")) + "\n"
        + json.dumps(_verdict("p-a", "7.3")) + "\n"
        + json.dumps({"provision_id": "p-err", "economy": "TL", "error": "ValidationError: x"}) + "\n",
        encoding="utf-8")
    (second / "map" / "verdicts_TL.jsonl").write_text(
        json.dumps(_verdict("p-both", "3.5")) + "\n"
        + json.dumps(_verdict("p-b", "12.9", in_force=False)) + "\n"
        + json.dumps({"provision_id": "p-err", "economy": "TL", "error": "ValidationError: x"}) + "\n",
        encoding="utf-8")

    h2 = tmp_path / "handoff2"
    (h2 / "source_text").mkdir(parents=True)
    with (h2 / "provisions.jsonl").open("w", encoding="utf-8") as f:
        for pid in ("p-both", "p-a", "p-b", "p-err"):
            f.write(json.dumps({"provision_id": pid, "economy": "TL", "doc_id": "d1",
                                "law_name": "An Act 2020", "article_section": "Art. 1",
                                "source_url": "https://x.gov/a.pdf",
                                "snippet_char_start": 0, "snippet_char_end": 5}) + "\n")
    (h2 / "source_text" / "d1.txt").write_text("Artigo 1\nshall be stored domestically\n",
                                               encoding="utf-8")

    monkeypatch.setattr(ex, "SETTINGS", dataclasses.replace(
        SETTINGS, out_dir=primary, handoff2_dir=h2, index_dir=tmp_path / "index"))
    monkeypatch.setattr(av, "SETTINGS", dataclasses.replace(SETTINGS, out_dir=primary))
    return primary, second


def _sheet(path, name):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb[name]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    hdr = [str(c) if c is not None else "" for c in rows[0]]
    data = [dict(zip(hdr, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]
    wb.close()
    return data


def test_the_workbook_carries_fires_from_every_arm(arms):
    primary, second = arms
    ex.export("TL", extra_dirs=[second])
    fires = _sheet(primary / "results" / "RDTII_P3_results_TL.xlsx", "All fires")
    assert sorted(f["Indicator"] for f in fires) == ["12.9", "3.5", "6.1", "7.3"], (
        "a one-arm export would show only two of these")


def test_a_provision_judged_in_both_arms_is_counted_once(arms):
    """`p-both` fires on 6.1 in one arm and 3.5 in the other. Two fires, ONE provision."""
    primary, second = arms
    ex.export("TL", extra_dirs=[second])
    summary = _sheet(primary / "results" / "RDTII_P3_results_TL.xlsx", "Summary")
    mapped = next(r["Value"] for r in summary if r["Metric"] == "Provisions mapped")
    assert mapped == 3, f"p-both, p-a, p-b — got {mapped}"


def test_an_error_row_present_in_both_arms_is_listed_once(arms):
    primary, second = arms
    ex.export("TL", extra_dirs=[second])
    errs = _sheet(primary / "results" / "RDTII_P3_results_TL.xlsx", "QA Errors")
    assert [e["provision_id"] for e in errs] == ["p-err"]


def test_a_not_in_force_provision_from_the_second_arm_reaches_the_sheet(arms):
    primary, second = arms
    ex.export("TL", extra_dirs=[second])
    nif = _sheet(primary / "results" / "RDTII_P3_results_TL.xlsx", "QA NotInForce")
    assert [n["provision_id"] for n in nif] == ["p-b"], "trap flags from either arm must show"


def test_the_workbook_is_written_to_the_primary_arm_only(arms):
    primary, second = arms
    ex.export("TL", extra_dirs=[second])
    assert (primary / "results" / "RDTII_P3_results_TL.xlsx").exists()
    assert not (second / "results" / "RDTII_P3_results_TL.xlsx").exists(), (
        "one file per country is the point")


def test_the_audit_page_unions_the_arms_too(arms):
    primary, second = arms
    av.render("TL", extra_dirs=[second])
    html = (primary / "audit" / "index_TL.html").read_text(encoding="utf-8")
    for ind in ("6.1", "7.3", "3.5", "12.9"):
        assert f'data-ind="{ind}"' in html, f"{ind} missing from the merged page"
    assert html.count("<tr data-ind=") == 4


def test_one_arm_still_works_unchanged(arms):
    """The extra-arm path must not be required. Five economies have exactly one arm."""
    primary, _ = arms
    ex.export("TL")
    fires = _sheet(primary / "results" / "RDTII_P3_results_TL.xlsx", "All fires")
    assert sorted(f["Indicator"] for f in fires) == ["6.1", "7.3"]


def test_a_missing_second_arm_is_skipped_not_fatal(arms, tmp_path):
    primary, _ = arms
    ex.export("TL", extra_dirs=[tmp_path / "does_not_exist"])
    fires = _sheet(primary / "results" / "RDTII_P3_results_TL.xlsx", "All fires")
    assert sorted(f["Indicator"] for f in fires) == ["6.1", "7.3"]
