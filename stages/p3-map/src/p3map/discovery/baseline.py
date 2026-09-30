r"""Host-baseline loader (KICKOFF build-order step 2).

Parses the COUNTRY sheets — never Consolidated, never Methodology — of the host databases into
out/baseline_rows.jsonl: the Round 1 book for AU/MY/SG, the Round 2 book for CN/LA (and the
other 2025 economies, which this stage does not map). Timor-Leste appears in neither, so every
TL fire is NEW by construction and the CSV Notes have to say so.

Used ONLY to tag a mapped row KNOWN or NEW after the fact, and to cite a governing law on a
no-provision row. It never chooses a source, a seed or a target: retrieval stays organic
(decision #11).

Sheet quirks handled (inspected 2026-07-15, re-inspected 2026-09-27 against the Round 2 book,
whose column layout is identical):
- pillar header rows carry the pillar NAME in Indicator_ID -> skipped
  (rows kept only when Indicator_ID matches r"\d+\.\d+", three-part IDs tolerated);
- merged Indicator_ID/Pillar_ID cells read as None on continuation rows ->
  forward-filled (the multi-row indicators 7.3/7.5);
- \xa0 (nbsp) cells are empty; strings whitespace-collapsed;
- references live from column H onward, across several unlabelled URL columns;
- the "Note" column is L on Singapore and China but M on Australia and Lao PDR, so it is
  located from each sheet's own header row instead of assumed. Round 1 read L unconditionally,
  which silently dropped every Australian note — and a note is where the baseline often cites
  the section that decides KNOWN vs NEW;
- indicator IDs are kept exactly as the host writes them ("6.1"). Round 1 rebuilt them into
  "P6-I1" here, which after the instrument hand-off would put every baseline row in a cell no
  fire can reach — NEW/KNOWN would tag everything NEW and report success.

`baseline_id` is load-bearing: Round 1's filed rows cite it in `baseline_match` and the reuse
path joins on it, so the Round 1 book keeps its "r1-<econ>-<sheet row>" ids unchanged. Round 2
rows carry an "r2-" prefix.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

from config.settings import SETTINGS

# Sheet name (lowercased, whitespace-collapsed) -> economy code. Both books, plus Timor-Leste
# for the day a host sheet covers it.
SHEET_ECON = {
    "singapore": "SG", "australia": "AU", "malaysia": "MY",            # Round 1, 2024
    "china": "CN", "india": "IN", "indonesia": "ID", "lao pdr": "LA",  # Round 2, 2025
    "mongolia": "MN", "russian federation": "RU", "thailand": "TH",
    # named in the sealed-draw list but absent from the Round 2 workbook we hold; mapped so that
    # a sheet arriving later needs no code change
    "viet nam": "VN", "vietnam": "VN", "kazakhstan": "KZ",
    "timor-leste": "TL", "timor leste": "TL",
}
# Consolidated duplicates the country sheets row for row; Methodology is the codebook.
SKIP_SHEETS = ("methodology", "consolidated")
IND_RE = re.compile(r"^(\d+)\.(\d+(?:\.\d+)?)$")
URL_RE = re.compile(r"https?://\S+")
FIRST_REF_COL = 7   # column H
MAX_COLS = 16


def _clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def _books() -> list[tuple[Path, str]]:
    """(workbook, baseline-id prefix), skipping any book this checkout does not carry."""
    books, missing = [], []
    for raw, prefix in ((SETTINGS.baseline_path, "r1"),
                        (SETTINGS.baseline_r2_path, "r2")):
        p = Path(raw)
        (books if p.exists() else missing).append((p, prefix))
    for p, _ in missing:
        print(f"[baseline] absent, skipped: {p}")
    if not books:
        raise SystemExit("[baseline] no host database found; set BASELINE_PATH / BASELINE_R2_PATH")
    return books


def _note_col(header: list[str]) -> int | None:
    for j, h in enumerate(header):
        if j >= FIRST_REF_COL and h.lower().startswith("note"):
            return j
    return None


def _sheet_rows(ws, sheet: str, econ: str, prefix: str) -> list[dict]:
    rows: list[dict] = []
    note_col: int | None = None
    last_pillar, last_ind = "", ""
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        cells = [_clean(c) for c in (list(r) + [""] * MAX_COLS)[:MAX_COLS]]
        if i == 1:
            note_col = _note_col(cells)
            continue
        pillar, ind_raw, score = cells[0], cells[1], cells[2]
        # forward-fill merged identity cells
        pillar = pillar or last_pillar
        carried = not IND_RE.match(cells[1] or "") and any(cells[2:6])
        ind_raw = ind_raw or (last_ind if carried else "")
        m = IND_RE.match(ind_raw)
        if not m:
            # pillar header row (name in Indicator_ID) or blank spacer
            if cells[1] and not IND_RE.match(cells[1]):
                last_ind = ""
            last_pillar = pillar or last_pillar
            continue
        last_pillar, last_ind = pillar, ind_raw
        p = m.group(1)
        note = cells[note_col] if note_col is not None else ""
        # Every column from H on, the note included: a note sometimes carries the only link
        # for a row, and Round 1's filed no-provision citations were built from this list.
        urls: list[str] = []
        for j in range(FIRST_REF_COL, MAX_COLS):
            urls += URL_RE.findall(cells[j])
        rows.append({
            "baseline_id": f"{prefix}-{econ.lower()}-{i:03d}",
            "round": prefix,
            "economy": econ,
            "sheet": sheet,
            "sheet_row": i,
            "pillar_id": p,
            "indicator_raw": ind_raw,
            "indicator": ind_raw,   # decimal text, the form every stage now uses
            "raw_score": score,
            "law": cells[3],
            "coverage": cells[4],
            "impact": cells[5],
            "timeframe": cells[6],
            "urls": urls,
            "note": note,
            "in_p3_scope": p in ("6", "7"),
        })
    return rows


def load_baseline() -> list[dict]:
    rows: list[dict] = []
    for path, prefix in _books():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            key = _clean(sheet).lower()
            if any(s in key for s in SKIP_SHEETS):
                continue
            econ = SHEET_ECON.get(key)
            if econ is None:
                print(f"[baseline] {path.name}: sheet {sheet!r} maps to no economy code, skipped")
                continue
            rows += _sheet_rows(wb[sheet], sheet, econ, prefix)
        wb.close()
    return rows


def run_baseline() -> None:
    rows = load_baseline()
    out = SETTINGS.out_dir / "baseline_rows.jsonl"
    SETTINGS.out_dir.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    scope = [r for r in rows if r["in_p3_scope"]]
    per: dict[str, int] = {}
    for r in scope:
        k = f"{r['economy']}:{r['indicator']}"
        per[k] = per.get(k, 0) + 1
    print(f"[baseline] {len(rows)} rows total, {len(scope)} in P6/P7 scope -> {out}")
    for k in sorted(per):
        print(f"  {k}: {per[k]}")


if __name__ == "__main__":
    run_baseline()
